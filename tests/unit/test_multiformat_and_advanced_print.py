import io
import json
import pytest
import PIL.Image
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from sqlalchemy import select

from backend.core.config import settings
from backend.models.configuration import PrintConfiguration
from backend.models.file import OrderFile
from backend.models.order import Order
from backend.models.user import User
from backend.schemas.llm import PageRule
from backend.services.files.converter import convert_and_validate_document
from backend.services.files.service import process_and_save_order_file
from backend.services.llm.parser import mock_llm_interpret
from backend.services.pricing.engine import calculate_order_pricing_snapshot
from backend.services.pricing.rules import validate_and_normalize_file_rules


# ---------------------------------------------------------------------------
# Requirement 3: Multi-Format Document Conversion & Page Counting Tests
# ---------------------------------------------------------------------------

def test_image_conversion_and_page_count():
    """Verifies image (JPEG / PNG) converts to 1-page standard printable PDF."""
    img = PIL.Image.new("RGB", (600, 800), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    jpeg_bytes = buf.getvalue()

    pdf_bytes, page_count, fmt = convert_and_validate_document(
        file_bytes=jpeg_bytes,
        original_filename="receipt.jpg",
        mime_type="image/jpeg",
    )
    assert fmt == "image"
    assert page_count == 1
    assert pdf_bytes.startswith(b"%PDF-")


def test_docx_conversion_and_page_count():
    """Verifies Microsoft Word (.docx) document converts into printable PDF."""
    doc = Document()
    doc.add_heading("Lab Assignment 1", level=1)
    doc.add_paragraph("This is the introductory problem statement for photocopy automation.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Param"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Pages"
    table.cell(1, 1).text = "10"
    table.cell(2, 0).text = "Sides"
    table.cell(2, 1).text = "Double"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    pdf_bytes, page_count, fmt = convert_and_validate_document(
        file_bytes=docx_bytes,
        original_filename="Assignment.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert fmt == "docx"
    assert page_count >= 1
    assert pdf_bytes.startswith(b"%PDF-")


def test_pptx_conversion_and_page_count():
    """Verifies PowerPoint (.pptx) converts to landscape PDF where 1 slide = 1 page."""
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Slide 1: Overview"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Slide 2: Implementation"
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = "Slide 3: Conclusions"

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    pdf_bytes, page_count, fmt = convert_and_validate_document(
        file_bytes=pptx_bytes,
        original_filename="Final_Presentation.pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert fmt == "pptx"
    assert page_count == 3
    assert pdf_bytes.startswith(b"%PDF-")


def test_xlsx_conversion_and_page_count():
    """Verifies Excel (.xlsx) spreadsheet renders tables into printable PDF."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(["ID", "Name", "Rate", "Quantity"])
    for i in range(1, 10):
        ws.append([i, f"Item {i}", 15.0 * i, i * 2])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    pdf_bytes, page_count, fmt = convert_and_validate_document(
        file_bytes=xlsx_bytes,
        original_filename="Balance_Sheet.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert fmt == "xlsx"
    assert page_count >= 1
    assert pdf_bytes.startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_process_and_save_order_docx_file(db_session):
    """Verifies process_and_save_order_file works with .docx and initializes 1 copy default."""
    user = User(whatsapp_user_id="wa_multiformat_user", phone_number="+919876543210")
    db_session.add(user)
    await db_session.flush()

    order = Order(order_code="ORD-DOCX-01", user_id=user.id, status="CREATED")
    db_session.add(order)
    await db_session.flush()

    doc = Document()
    doc.add_paragraph("Unit test docx content for printing")
    buf = io.BytesIO()
    doc.save(buf)

    saved_file = await process_and_save_order_file(
        session=db_session,
        order=order,
        original_filename="Project_Synopsis.docx",
        file_bytes=buf.getvalue(),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert saved_file.original_filename == "Project_Synopsis.docx"
    assert saved_file.page_count >= 1

    cfg_res = await db_session.execute(
        select(PrintConfiguration).where(PrintConfiguration.file_id == saved_file.id)
    )
    cfg = cfg_res.scalar_one()
    assert cfg.copies == 1
    assert cfg.color_mode == "bw"
    assert cfg.default_sides == "single"


# ---------------------------------------------------------------------------
# Requirement 2: Multiple Copies, Duplex, & Partial Single/Double Tests
# ---------------------------------------------------------------------------

def test_nlp_multiple_copies_and_back_to_back():
    """Verifies NLP parser correctly extracts copy counts and back-to-back / duplex."""
    dummy_files = [
        OrderFile(upload_sequence=1, original_filename="thesis.pdf", page_count=20),
    ]
    res = mock_llm_interpret("Print 3 copies double sided back to back", dummy_files)
    assert not res.ambiguity
    assert len(res.file_rules) == 1
    rule = res.file_rules[0]
    assert rule.copies == 3
    assert rule.default_sides == "double"
    assert rule.color_mode == "bw"


def test_nlp_partial_single_and_double_sided_explicit_ranges():
    """Verifies NLP parser handles partial single and double sided page ranges."""
    dummy_files = [
        OrderFile(upload_sequence=1, original_filename="book.pdf", page_count=30),
    ]
    # Pages 1-5 single sided, pages 6-30 double sided
    res = mock_llm_interpret("Pages 1 to 5 single sided and pages 6 to 30 double sided", dummy_files)
    assert not res.ambiguity
    assert len(res.file_rules) == 1
    rule = res.file_rules[0]
    assert len(rule.rules) == 2
    assert rule.rules[0].page_start == 1
    assert rule.rules[0].page_end == 5
    assert rule.rules[0].sides == "single"
    assert rule.rules[1].page_start == 6
    assert rule.rules[1].page_end == 30
    assert rule.rules[1].sides == "double"


def test_nlp_first_n_single_rest_double_with_copies():
    """Verifies NLP parser handles first N single, rest back-to-back with 2 copies."""
    dummy_files = [
        OrderFile(upload_sequence=1, original_filename="notes.pdf", page_count=10),
    ]
    res = mock_llm_interpret("First 2 pages single sided, rest double sided, 2 copies", dummy_files)
    assert not res.ambiguity
    assert len(res.file_rules) == 1
    rule = res.file_rules[0]
    assert rule.copies == 2
    assert len(rule.rules) == 2
    assert rule.rules[0].page_start == 1
    assert rule.rules[0].page_end == 2
    assert rule.rules[0].sides == "single"
    assert rule.rules[1].page_start == 3
    assert rule.rules[1].page_end == 10
    assert rule.rules[1].sides == "double"


def test_nlp_per_file_configurations():
    """Verifies NLP parser handles per-file specific sides and copies."""
    dummy_files = [
        OrderFile(upload_sequence=1, original_filename="form.pdf", page_count=2),
        OrderFile(upload_sequence=2, original_filename="handout.pdf", page_count=10),
    ]
    res = mock_llm_interpret("File 1 single sided 3 copies, File 2 double sided 1 copy", dummy_files)
    assert not res.ambiguity
    assert len(res.file_rules) == 2
    assert res.file_rules[0].file_reference == 1
    assert res.file_rules[0].copies == 3
    assert res.file_rules[0].default_sides == "single"

    assert res.file_rules[1].file_reference == 2
    assert res.file_rules[1].copies == 1
    assert res.file_rules[1].default_sides == "double"


@pytest.mark.asyncio
async def test_pricing_with_multiple_copies_and_partial_rules(db_session):
    """
    Verifies that pricing engine correctly multiplies page costs and page counts by copies,
    and applies partial single/double rates.
    """
    user = User(whatsapp_user_id="wa_pricing_copies", phone_number="+919988112233")
    db_session.add(user)
    await db_session.flush()

    order = Order(order_code="ORD-COPIES-PRICING", user_id=user.id, status="CREATED")
    db_session.add(order)
    await db_session.flush()

    # Document: 10 pages total
    # Pages 1-2: single-sided (2 pages * 2.0 = 4.0 per copy)
    # Pages 3-10: double-sided (8 pages * 1.5 = 12.0 per copy)
    # Total single copy cost = 16.0
    # 3 copies -> 16.0 * 3 = 48.0
    # Printed pages -> 10 * 3 = 30 pages
    # 1 file base charge -> 10.0
    # Total amount -> 58.0
    file1 = OrderFile(
        order_id=order.id,
        upload_sequence=1,
        original_filename="booklet.pdf",
        file_hash="hash_copies",
        mime_type="application/pdf",
        storage_path="/tmp/booklet.pdf",
        page_count=10,
        file_size=2048,
        status="VALIDATED",
    )
    db_session.add(file1)
    await db_session.flush()

    rules = [
        {"page_start": 1, "page_end": 2, "sides": "single"},
        {"page_start": 3, "page_end": 10, "sides": "double"},
    ]
    config1 = PrintConfiguration(
        order_id=order.id,
        file_id=file1.id,
        copies=3,
        color_mode="bw",
        default_sides="single",
        rules_json=json.dumps(rules),
    )
    db_session.add(config1)
    await db_session.flush()

    snapshot = await calculate_order_pricing_snapshot(db_session, order)

    expected_base = round(1 * settings.PRICE_FILE_BASE_CHARGE, 2)
    single_copy_cost = 2 * settings.PRICE_BW_SINGLE + 8 * settings.PRICE_BW_DOUBLE
    expected_printing = round(single_copy_cost * 3, 2)
    expected_total = round(expected_base + expected_printing, 2)

    assert snapshot.file_count == 1
    assert snapshot.total_pages == 30  # 10 pages * 3 copies
    assert snapshot.file_base_charge == expected_base
    assert snapshot.printing_charge == expected_printing
    assert snapshot.total_amount == expected_total
    assert order.total_amount == expected_total
    assert order.total_pages == 30
    assert len(snapshot.details) == 1
    assert snapshot.details[0].copies == 3
    assert snapshot.details[0].bw_single_pages == 6   # 2 * 3
    assert snapshot.details[0].bw_double_pages == 24  # 8 * 3
    assert snapshot.details[0].file_cost == expected_printing


def test_nlp_exact_user_scenarios_multi_document_copies():
    """
    Tests exact customer instructions:
    1. '4 copies of 1 st document and single side other'
    2. 'file 1  - 4 copies \n rest file single cop'
    3. 'no 4 copies of first document and single copies of all rest'
    """
    files = [
        OrderFile(upload_sequence=1, original_filename="doc1.pdf", page_count=5),
        OrderFile(upload_sequence=2, original_filename="doc2.pdf", page_count=8),
        OrderFile(upload_sequence=3, original_filename="doc3.pdf", page_count=2),
    ]

    # Case 1: "4 copies of 1 st document and single side other"
    res1 = mock_llm_interpret("4 copies of 1 st document and single side other", files)
    assert not res1.ambiguity
    assert len(res1.file_rules) == 3
    assert res1.file_rules[0].copies == 4
    assert res1.file_rules[0].default_sides == "single"
    assert res1.file_rules[1].copies == 1
    assert res1.file_rules[1].default_sides == "single"
    assert res1.file_rules[2].copies == 1
    assert res1.file_rules[2].default_sides == "single"

    # Case 2: "file 1  - 4 copies \n rest file single cop"
    res2 = mock_llm_interpret("file 1  - 4 copies \n rest file single cop", files)
    assert not res2.ambiguity
    assert len(res2.file_rules) == 3
    assert res2.file_rules[0].copies == 4
    assert res2.file_rules[1].copies == 1
    assert res2.file_rules[2].copies == 1

    # Case 3: "no 4 copies of first document and single copies of all rest"
    res3 = mock_llm_interpret("no 4 copies of first document and single copies of all rest", files)
    assert not res3.ambiguity
    assert len(res3.file_rules) == 3
    assert res3.file_rules[0].copies == 4
    assert res3.file_rules[1].copies == 1
    assert res3.file_rules[2].copies == 1


@pytest.mark.asyncio
async def test_whatsapp_bot_short_acknowledgment_and_pricing_for_copies(db_session, monkeypatch):
    """
    Verifies that WhatsApp bot:
    1. Immediately acknowledges each file with a concise 'Total X received' line
    2. Correctly updates order configuration, copies, and pricing when user sends instructions like:
       '4 copies of 1 st document and single side other'
    """
    from backend.services.whatsapp.bot import get_active_customer_order, handle_inbound_whatsapp_event
    from backend.services.whatsapp.service import whatsapp_service

    sent_messages = []
    async def mock_send_text(to_phone: str, text: str) -> bool:
        sent_messages.append({"to": to_phone, "text": text})
        return True

    async def mock_download(info):
        return b"%PDF-1.4 mock pdf bytes"

    monkeypatch.setattr(whatsapp_service, "send_text_message", mock_send_text)
    monkeypatch.setattr(whatsapp_service, "download_media", mock_download)

    # Mock PyMuPDF / fitz to return 4 pages for doc1 and 6 pages for doc2
    call_count = 0
    class MockDoc:
        def __init__(self, *args, **kwargs):
            nonlocal call_count
            call_count += 1
            self.page_count = 4 if call_count == 1 else 6
            self.is_encrypted = False
        def __len__(self):
            return self.page_count
        def close(self):
            pass

    import fitz
    monkeypatch.setattr(fitz, "open", lambda *args, **kwargs: MockDoc())

    phone = "+919876543210"

    # 1. User sends first document
    event1 = {
        "type": "document",
        "sender": phone,
        "display_name": "Multi Doc User",
        "document": {"filename": "file1.pdf", "mime_type": "application/pdf"},
    }
    await handle_inbound_whatsapp_event(db_session, event1)

    assert len(sent_messages) == 1
    ack1 = sent_messages[-1]["text"]
    assert "Total 1 file received:" in ack1
    assert "file1.pdf" in ack1
    assert "File 1 of 10 Received!" in ack1

    # 2. User sends second document
    event2 = {
        "type": "document",
        "sender": phone,
        "display_name": "Multi Doc User",
        "document": {"filename": "file2.pdf", "mime_type": "application/pdf"},
    }
    await handle_inbound_whatsapp_event(db_session, event2)

    assert len(sent_messages) == 2
    ack2 = sent_messages[-1]["text"]
    assert "Total 2 files received:" in ack2
    assert "file2.pdf" in ack2
    assert "File 2 of 10 Received!" in ack2

    # 3. User sends print instructions: "4 copies of 1 st document and single side other"
    instruction_event = {
        "type": "text",
        "sender": phone,
        "display_name": "Multi Doc User",
        "text": "4 copies of 1 st document and single side other",
    }
    await handle_inbound_whatsapp_event(db_session, instruction_event)

    assert len(sent_messages) == 3
    cfg_reply = sent_messages[-1]["text"]
    assert "Print Options Configured" in cfg_reply
    assert "File Details" in cfg_reply
    assert "4 Copies" in cfg_reply
    assert "1 Copy" in cfg_reply

    # Verify order state in DB
    user_query = select(User).where(User.phone_number == phone)
    user = (await db_session.execute(user_query)).scalar_one()
    order = await get_active_customer_order(db_session, user)

    assert order is not None
    assert len(order.files) == 2
    file1 = order.files[0]
    file2 = order.files[1]

    assert file1.configuration.copies == 4
    assert file2.configuration.copies == 1
    # file1: 4 pages * 4 copies = 16 pages
    # file2: 6 pages * 1 copy = 6 pages
    # Total = 22 pages
    assert order.total_pages == 22
